import sys
import os
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, Color
)
from openpyxl.utils import get_column_letter

REGISTER_SIZE = 26
TAP_DEGREES = [26, 8, 7, 1]

DEFAULT_SEED = "1" * REGISTER_SIZE
DEFAULT_STEPS = 60


def validate_seed(seed: str) -> str:
    seed = seed.strip()
    if len(seed) != REGISTER_SIZE:
        raise ValueError(
            f"Начальное состояние должно содержать ровно {REGISTER_SIZE} бит, "
            f"получено: {len(seed)}"
        )
    if not all(c in "01" for c in seed):
        raise ValueError(
            "Начальное состояние должно содержать только символы 0 и 1")
    if all(c == "0" for c in seed):
        raise ValueError(
            "Начальное состояние не может быть нулевым (регистр зациклится на нулях)")
    return seed


def degree_to_col_index(degree: int) -> int:
    return 3 + (REGISTER_SIZE - degree)


COL_LABEL = 1  # A  — Steps / Register state
COL_STEP = 2  # B  — step number
COL_FIRST = 3  # C  — highest (m power)
COL_LAST = COL_FIRST + REGISTER_SIZE - 1  # AB — lowest (1 power)
COL_XOR = COL_LAST + 1  # AC — XOR

TAP_COLS = {degree_to_col_index(d) for d in TAP_DEGREES}


def _thin_border() -> Border:
    s = Side(border_style="thin", color="000000")
    return Border(left=s, right=s, top=s, bottom=s)


def _fill_rgb(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


FILL_TAP = _fill_rgb("4472C4")  # blue — taps
FILL_XOR = _fill_rgb("ED7D31")  # orange — XOR column
FILL_NONE = PatternFill(fill_type=None)

BORDER = _thin_border()
FONT_BASE = Font(name="Calibri", size=11)
FONT_BOLD = Font(name="Calibri", size=11, bold=True)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left",   vertical="center")


def _style(cell, fill=None, font=None, alignment=None, border=True):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = BORDER


def xor_formula(row: int) -> str:
    args = ", ".join(
        f"{get_column_letter(degree_to_col_index(d))}{row}"
        for d in TAP_DEGREES
    )
    return f"=IF(_xlfn.XOR({args}),1,0)"


def generate(seed: str, n_steps: int, output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "LFSR"

    DATA_ROW_START = 3  # first row with data (initial state)
    DATA_ROW_END = DATA_ROW_START + n_steps - 1  # last row

    # column width
    ws.column_dimensions[get_column_letter(COL_LABEL)].width = 3.0  # A
    ws.column_dimensions[get_column_letter(COL_STEP)].width = 3.0  # B

    for col_idx in range(COL_FIRST, COL_XOR + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 5.0

    # row height
    ws.row_dimensions[1].height = 17.0
    ws.row_dimensions[2].height = 15.0
    for row in range(DATA_ROW_START, DATA_ROW_END + 1):
        ws.row_dimensions[row].height = 14.5

    # row 1: labels (powers)
    label_start = get_column_letter(COL_FIRST + 1)
    label_end = get_column_letter(COL_LAST)
    ws.merge_cells(f"{label_start}1:{label_end}1")
    cell_degrees = ws[f"{label_start}1"]
    cell_degrees.value = "Степени"
    cell_degrees.alignment = ALIGN_CENTER
    cell_degrees.font = FONT_BASE

    # XOR
    xor_col_letter = get_column_letter(COL_XOR)
    cell_xor_hdr = ws[f"{xor_col_letter}1"]
    cell_xor_hdr.value = "xor"
    cell_xor_hdr.alignment = ALIGN_CENTER
    cell_xor_hdr.font = FONT_BASE
    _style(cell_xor_hdr, fill=FILL_XOR, border=True)

    # row 2: power numbers
    for degree in range(REGISTER_SIZE, 0, -1):
        col_idx = degree_to_col_index(degree)
        col_ltr = get_column_letter(col_idx)
        cell = ws[f"{col_ltr}2"]
        cell.value = degree
        cell.alignment = ALIGN_CENTER
        cell.font = FONT_BASE
        cell.border = BORDER
        if col_idx in TAP_COLS:
            cell.fill = FILL_TAP

    # row 3: initial state + xor
    for degree in range(REGISTER_SIZE, 0, -1):
        col_idx = degree_to_col_index(degree)
        col_ltr = get_column_letter(col_idx)
        bit_pos = REGISTER_SIZE - degree
        cell = ws[f"{col_ltr}{DATA_ROW_START}"]
        cell.value = int(seed[bit_pos])
        cell.alignment = ALIGN_CENTER
        cell.font = FONT_BASE
        cell.border = BORDER
        if col_idx in TAP_COLS:
            cell.fill = FILL_TAP

    cell_xor0 = ws[f"{xor_col_letter}{DATA_ROW_START}"]
    cell_xor0.value = xor_formula(DATA_ROW_START)
    cell_xor0.alignment = ALIGN_CENTER
    cell_xor0.font = FONT_BASE
    _style(cell_xor0, fill=FILL_XOR, border=True)

    #  rows 4…(n_steps+2): shifts via Excel-formulas
    for row in range(DATA_ROW_START + 1, DATA_ROW_END + 1):
        prev = row - 1

        for degree in range(REGISTER_SIZE, 0, -1):
            col_idx = degree_to_col_index(degree)
            col_ltr = get_column_letter(col_idx)
            cell = ws[f"{col_ltr}{row}"]
            cell.alignment = ALIGN_CENTER
            cell.font = FONT_BASE
            cell.border = BORDER
            if col_idx in TAP_COLS:
                cell.fill = FILL_TAP

            if degree == 1:
                cell.value = f"={xor_col_letter}{prev}"
            else:
                src_col = get_column_letter(degree_to_col_index(degree - 1))
                cell.value = f"={src_col}{prev}"

        cell_xor = ws[f"{xor_col_letter}{row}"]
        cell_xor.value = xor_formula(row)
        cell_xor.alignment = ALIGN_CENTER
        cell_xor.font = FONT_BASE
        _style(cell_xor, fill=FILL_XOR, border=True)

    # Columns A (label) and B (step number)
    ws.merge_cells(f"A{DATA_ROW_START}:A{DATA_ROW_END}")
    cell_label = ws[f"A{DATA_ROW_START}"]
    cell_label.value = "Шаги / Состояние регистра"
    cell_label.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True)
    cell_label.font = FONT_BASE
    cell_label.border = BORDER

    for step, row in enumerate(range(DATA_ROW_START, DATA_ROW_END + 1), start=1):
        cell_step = ws[f"B{row}"]
        cell_step.value = step
        cell_step.alignment = ALIGN_CENTER
        cell_step.font = FONT_BASE
        cell_step.border = BORDER

    for col_ltr in ("A", "B"):
        for r in (1, 2):
            ws[f"{col_ltr}{r}"].border = BORDER

    ws.freeze_panes = "C3"

    wb.save(output_path)
    print(f"Файл сохранён: {output_path}")
    print(f"  Многочлен:          P(x) = x^26 + x^8 + x^7 + x + 1")
    print(f"  Начальное состояние: {seed}")
    print(f"  Количество шагов:   {n_steps}")
    print(f"  Отводы (степени):   {TAP_DEGREES}")


def main():
    seed = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SEED
    n_steps = int(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_STEPS

    try:
        seed = validate_seed(seed)
    except ValueError as e:
        print(f"Ошибка: {e}", file=sys.stderr)
        sys.exit(1)

    if n_steps < 1:
        print("Ошибка: количество шагов должно быть не менее 1", file=sys.stderr)
        sys.exit(1)

    output_name = f"LFSR_m{REGISTER_SIZE}_steps{n_steps}.xlsx"
    generate(seed, n_steps, output_name)


if __name__ == "__main__":
    main()
